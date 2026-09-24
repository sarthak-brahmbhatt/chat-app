import hashlib
import uuid
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph


@dataclass(frozen=True)
class ReferenceChunk:
    point_id: str
    text: str
    payload: dict


def iter_blocks(document: DocumentObject):
    """Yield paragraphs and tables in their original document order."""
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def table_text(table: Table) -> str:
    rows = []
    for row in table.rows:
        rows.append(" | ".join(cell.text.strip().replace("\n", " / ") for cell in row.cells))
    return "\n".join(rows)


def section_text(path: Path, heading: str) -> str:
    document = Document(path)
    collecting = False
    lines: list[str] = []

    for block in iter_blocks(document):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            level = _heading_level(block)
            if text == heading:
                collecting = True
                lines.append(text)
                continue
            if collecting and level is not None and level <= 2:
                break
            if collecting and text:
                lines.append(text)
        elif collecting:
            lines.append(table_text(block))

    if not lines:
        raise ValueError(f"Section '{heading}' was not found in {path.name}")
    return "\n".join(lines)


def coding_chunks(path: Path, version: str = "FY2024") -> list[ReferenceChunk]:
    document = Document(path)
    code_type: str | None = None
    category: str | None = None
    chunks: list[ReferenceChunk] = []

    for block in iter_blocks(document):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if text.startswith("ICD-10-CM"):
                code_type = "ICD-10"
                category = None
            elif text.startswith("CPT Codes"):
                code_type = "CPT"
                category = None
            elif _heading_level(block) == 2:
                code_type = None
            elif code_type and _heading_level(block) == 3:
                category = text
            continue

        if not code_type or not category:
            continue

        headers = [cell.text.strip().lower() for cell in block.rows[0].cells]
        for row in block.rows[1:]:
            values = [cell.text.strip() for cell in row.cells]
            record = dict(zip(headers, values, strict=False))
            code = record.get("code", "")
            description = record.get("description", "")
            notes = record.get("clinical notes") or record.get("usage") or ""
            text = f"{code_type} code {code}: {description}. Category: {category}. Notes: {notes}"
            chunks.append(_chunk(
                key=f"{code_type}:{code}",
                text=text,
                payload={
                    "chunk_type": "coding",
                    "code_type": code_type,
                    "code": code,
                    "description": description,
                    "category": category,
                    "clinical_notes": notes,
                    "source_document": path.name,
                    "source_section": category,
                    "version": version,
                },
            ))
    return chunks


def quality_measure_chunks(path: Path, version: str = "provided-spec") -> list[ReferenceChunk]:
    document = Document(path)
    chunks: list[ReferenceChunk] = []
    current_heading: str | None = None
    current_lines: list[str] = []

    def finish_measure() -> None:
        if not current_heading:
            return
        text = "\n".join(current_lines)
        measure_id = _measure_id(current_heading)
        chunks.append(_chunk(
            key=f"measure:{measure_id}",
            text=text,
            payload={
                "chunk_type": "quality_measure",
                "measure_id": measure_id,
                "measure_name": current_heading,
                "source_document": path.name,
                "source_section": current_heading,
                "version": version,
            },
        ))

    for block in iter_blocks(document):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if _heading_level(block) == 2:
                finish_measure()
                current_heading = text if text.startswith("Measure ") else None
                current_lines = [text] if current_heading else []
            elif current_heading and text:
                current_lines.append(text)
        elif current_heading:
            current_lines.append(table_text(block))

    finish_measure()
    return chunks


def instruction_sections(resources_dir: Path) -> dict[str, str]:
    source = resources_dir / "source"
    coding = source / "coding-reference.docx"
    quality = source / "quality-measures-spec.docx"
    return {
        "coding_guidelines": section_text(coding, "Coding Guidelines — Key Principles"),
        "tool_call_specification": section_text(coding, "Tool Call Specification"),
        "general_extraction_guidelines": section_text(quality, "General Extraction Guidelines"),
    }


def _heading_level(paragraph: Paragraph) -> int | None:
    name = paragraph.style.name
    if not name.startswith("Heading "):
        return None
    try:
        return int(name.split()[-1])
    except ValueError:
        return None


def _measure_id(heading: str) -> str:
    # "Measure 1: SEP-1 — ..." -> "SEP-1"
    after_colon = heading.split(":", 1)[1].strip()
    return after_colon.split("—", 1)[0].strip()


def _chunk(key: str, text: str, payload: dict) -> ReferenceChunk:
    checksum = hashlib.sha256(text.encode("utf-8")).hexdigest()
    point_id = str(uuid.uuid5(uuid.NAMESPACE_URL, f"clinical-reference:{key}"))
    return ReferenceChunk(point_id, text, {**payload, "text": text, "checksum": checksum})
